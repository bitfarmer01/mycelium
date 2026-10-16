"""SurrealDB-only corpus path: data/docs/**/*.{md,txt,html,pdf,csv,tsv,xlsx} -> chunk + source rows (bge-large-en-v1.5, 1024-d); data/questions.json -> question rows.

Idempotent: chunk ids are sha1(abs path + index), source ids sha1(abs path), question ids from the file.
"""
from __future__ import annotations

import asyncio
import datetime as _dt
import hashlib
import html as _html
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core import db  # noqa: E402
from core.contracts import CONFIG  # noqa: E402

DOCS = ROOT / "data" / "docs"
QUESTIONS = ROOT / "data" / "questions.json"
MODEL = "BAAI/bge-large-en-v1.5"
EXTS = {".md", ".txt", ".html", ".htm", ".pdf", ".csv", ".tsv", ".xlsx"}
SHEET_EXTS = {".csv", ".tsv", ".xlsx"}
# Hard limits (apply to CLI ingest and to /upload): keep one file from stalling the embedder or flooding the corpus.
PDF_MAX_PAGES = 300
SHEET_MAX_SHEETS = 20
SHEET_MAX_ROWS = 5000  # per sheet, header included
SHEET_MAX_COLS = 50
ROWS_PER_PARA = 20  # spreadsheet rows per paragraph, so chunks break between row groups
LINKS = DOCS / "sources.json"  # optional {"relative/path.pdf": "https://..."} for files that cannot carry a source_link line
EMBED_BATCH = 32
INSERT_BATCH = 200


# ---------- pure helpers (unit-tested) ----------

def chunk_text(text: str, words: int = 350, overlap: int = 50) -> list[str]:
    """Split on paragraph boundaries into chunks of <= `words` words; consecutive chunks share exactly `overlap` words.

    A chunk ends at the last paragraph boundary that fits; if none fits (huge paragraph) it hard-cuts at `words`.
    """
    paras = [p.split() for p in re.split(r"\n\s*\n", text) if p.split()]
    toks: list[str] = []
    bounds: set[int] = set()  # token index where a new paragraph starts
    for p in paras:
        if toks:
            bounds.add(len(toks))
        toks.extend(p)
    total = len(toks)
    if total == 0:
        return []
    out: list[str] = []
    start = 0
    while True:
        limit = min(start + words, total)
        end = limit
        if limit < total:
            cands = [b for b in bounds if start + overlap < b <= limit]
            if cands:
                end = max(cands)
        parts, cur = [], [toks[start]]
        for i in range(start + 1, end):
            if i in bounds:
                parts.append(" ".join(cur))
                cur = []
            cur.append(toks[i])
        parts.append(" ".join(cur))
        out.append("\n\n".join(parts))
        if end >= total:
            return out
        start = end - overlap


def strip_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr|section|article)>", "\n\n", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    raw = _html.unescape(raw)
    return re.sub(r"[ \t]+", " ", raw)


def extract_title(text: str, filename: str) -> str:
    m = re.search(r"(?m)^#\s+(.+?)\s*#*\s*$", text)
    return m.group(1).strip() if m else Path(filename).stem


def html_title(raw: str) -> str | None:
    m = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw) or re.search(r"(?is)<h1[^>]*>(.*?)</h1>", raw)
    return _html.unescape(re.sub(r"<[^>]+>", "", m.group(1))).strip() or None if m else None


def front_matter_link(text: str) -> tuple[str | None, str]:
    """Return (source_link, text_without_that_line) when a `source_link:` line is present near the top."""
    m = re.search(r"(?mi)^\s*source_link:\s*(\S+)\s*$", text[:2000])
    if not m:
        return None, text
    return m.group(1), text[: m.start()] + text[m.end():]


def chunk_id(path: str, index: int) -> str:
    return hashlib.sha1(f"{path}#{index}".encode()).hexdigest()[:20]


def source_id(path: str) -> str:
    return hashlib.sha1(path.encode()).hexdigest()[:20]


def rows_text(rows: list[list]) -> str:
    """Spreadsheet rows -> 'Header: value; Header: value' lines, first row as header, blank line every ROWS_PER_PARA rows."""
    rows = [[("" if c is None else str(c).strip()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    head, lines = rows[0], []
    for n, r in enumerate(rows[1:], 1):
        cells = [f"{head[i] if i < len(head) and head[i] else f'col{i + 1}'}: {v}" for i, v in enumerate(r) if v]
        lines.append("; ".join(cells))
        if n % ROWS_PER_PARA == 0:
            lines.append("")
    return "\n".join(lines)


def sheet_parts(path: Path) -> list[dict]:
    """One part per sheet (csv/tsv are a single sheet)."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return [{"label": ws.title, "meta": {"sheet": ws.title},
                     "text": rows_text([list(r)[:SHEET_MAX_COLS] for r in ws.iter_rows(max_row=SHEET_MAX_ROWS, values_only=True)])}
                    for ws in wb.worksheets[:SHEET_MAX_SHEETS]]
        finally:
            wb.close()
    import csv

    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(encoding="utf-8", errors="replace", newline="") as fh:
        rows = [r[:SHEET_MAX_COLS] for _, r in zip(range(SHEET_MAX_ROWS), csv.reader(fh, delimiter=delim))]
        return [{"label": None, "text": rows_text(rows), "meta": {}}]


def pdf_parts(path: Path) -> tuple[str | None, list[dict]]:
    """(metadata title, one part per page with text)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = None
    try:
        title = (reader.metadata.title or "").strip() or None if reader.metadata else None
    except Exception:
        pass
    parts = []
    if reader.is_encrypted:
        raise ValueError("encrypted PDF")
    for n, page in enumerate(reader.pages[:PDF_MAX_PAGES], 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append({"label": f"p. {n}", "text": text, "meta": {"page": n}})
    return title, parts


def manifest_links(root: Path | None = None) -> dict[str, str]:
    f = (root / "sources.json") if root else LINKS
    try:
        return json.loads(f.read_text()) if f.exists() else {}
    except Exception:
        return {}


def load_doc(path: Path, links: dict[str, str] | None = None) -> dict:
    """{path, title, link, timestamp, parts:[{label, text, meta}]}; pdf -> one part per page, spreadsheet -> one per sheet."""
    ext = path.suffix.lower()
    ap = str(path.resolve())
    ts = _dt.datetime.fromtimestamp(path.stat().st_mtime, _dt.timezone.utc).isoformat()
    links = links or {}
    link = links.get(str(path.relative_to(DOCS)) if path.is_relative_to(DOCS) else path.name) or links.get(path.name)
    if ext == ".pdf":
        title, parts = pdf_parts(path)
        title = title or path.stem.replace("_", " ").replace("-", " ")
    elif ext in SHEET_EXTS:
        parts = sheet_parts(path)
        title = path.stem.replace("_", " ").replace("-", " ")
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
        fm_link, raw = front_matter_link(raw)
        link = fm_link or link
        if ext in (".html", ".htm"):
            title = html_title(raw) or path.stem
            text = strip_html(raw)
        else:
            text = re.sub(r"(?s)\A---\n.*?\n---\n", "", raw)  # drop YAML front matter
            title = extract_title(text, path.name)
        parts = [{"label": None, "text": text, "meta": {}}]
    return {"path": ap, "title": title, "link": link or f"file://{ap}", "timestamp": ts, "parts": parts, "kind": ext.lstrip(".")}


def doc_files(root: Path = DOCS) -> list[Path]:
    if not root.exists():
        return []
    return sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in EXTS)


# ---------- embedding + DB ----------

_model = None


def embed(texts: list[str]) -> list[list[float]]:
    global _model
    if _model is None:
        try:
            from core.retrieve import _get_model  # same bge-large instance the query path uses

            _model = _get_model()
        except ImportError:
            from sentence_transformers import SentenceTransformer  # lazy: unit tests never load the model

            _model = SentenceTransformer(MODEL)
    model = _model
    out: list[list[float]] = []
    for i in range(0, len(texts), EMBED_BATCH):
        vecs = model.encode(texts[i:i + EMBED_BATCH], normalize_embeddings=True)
        out.extend([float(x) for x in v] for v in vecs)
    return out


async def write_chunks(rows: list[dict]) -> None:
    sql = (
        "INSERT INTO chunk $rows ON DUPLICATE KEY UPDATE "
        "document = $input.document, title = $input.title, source_link = $input.source_link, "
        "metadata = $input.metadata, timestamp = $input.timestamp, embedding = $input.embedding;"
    )
    for i in range(0, len(rows), INSERT_BATCH):
        await db.q(sql, {"rows": rows[i:i + INSERT_BATCH]})


class TooManyChunks(ValueError):
    pass


def build_chunks(files: list[Path], links: dict[str, str] | None = None,
                 max_chunks: int | None = None) -> tuple[list[dict], list[dict]]:
    """Sync part of ingest (load, chunk, embed): safe to run in a worker thread. Returns (chunk rows with embeddings, source rows)."""
    links = manifest_links() if links is None else links
    chunks: list[dict] = []
    sources: list[dict] = []
    for f in files:
        try:
            d = load_doc(f, links)
        except Exception as e:  # one unreadable file must not stop the corpus
            print(f"SKIP {f.name}: {type(e).__name__}")
            continue
        sid = source_id(d["path"])
        sources.append({"id": sid, "title": d["title"]})
        i = 0
        for part in d["parts"]:
            title = f"{d['title']} ({part['label']})" if part["label"] else d["title"]
            for c in chunk_text(part["text"]):
                chunks.append({
                    "id": chunk_id(d["path"], i),
                    "document": c,
                    "title": title,
                    "source_link": d["link"],
                    "metadata": {"path": d["path"], "index": i, "source": f"source:{sid}", "words": len(c.split()),
                                 "kind": d["kind"], **part["meta"]},
                    "timestamp": d["timestamp"],
                    "alpha": 1, "beta": 1, "cited_count": 0, "retrieved_count": 0, "last_event": None,
                })
                i += 1
    if max_chunks is not None and len(chunks) > max_chunks:  # checked before the expensive embed
        raise TooManyChunks(f"{len(chunks)} chunks exceeds the limit of {max_chunks}")
    if chunks:
        vecs = embed([c["document"] for c in chunks])
        for c, v in zip(chunks, vecs):
            c["embedding"] = v
    return chunks, sources


async def store(chunks: list[dict], sources: list[dict], replace_paths: list[str] | None = None) -> None:
    """Write chunks and sources. replace_paths: drop older chunks of those files first, so a re-upload never leaves stale text."""
    for path in replace_paths or []:
        await db.q("DELETE chunk WHERE metadata.path = $p;", {"p": path})
    if chunks:
        await write_chunks(chunks)
    for s in sources:
        await db.q(
            "UPSERT type::record('source', $id) MERGE { title: $title }; "
            "UPDATE type::record('source', $id) SET alpha = alpha ?? 1, beta = beta ?? 1;",
            s,
        )


async def ingest_docs(files: list[Path] | None = None) -> tuple[int, int]:
    chunks, sources = build_chunks(doc_files() if files is None else files)
    await store(chunks, sources)
    return len(chunks), len(sources)


async def ingest_questions() -> int:
    if not QUESTIONS.exists():
        print("questions=0 NEED data/questions.json")
        return 0
    from scripts.make_split import frozen_split, load_questions

    questions = load_questions(QUESTIONS)
    heldout, _, _ = frozen_split(questions, CONFIG["held_out_n"], CONFIG["seed"])
    held = set(heldout)
    for qd in questions:
        qid = str(qd["id"])
        row = {
            "text": qd.get("text") or qd.get("question", ""),
            "answer": qd.get("answer"),
            "shape": qd.get("shape"),
            "session_id": qd.get("session_id"),
            "readings": qd.get("readings", []),
            "taken": qd.get("taken"),
            "split": "heldout" if qid in held else "train",
            "ts": qd.get("ts") or db.now(),
        }
        await db.q("UPSERT type::record('question', $id) MERGE $row;", {"id": qid, "row": row})
    return len(questions)


async def main() -> None:
    n_chunk, n_source = await ingest_docs()
    if n_chunk == 0:
        print(f"NEED documents in {DOCS} (md, txt, html, pdf, csv, tsv, xlsx)")
    n_q = await ingest_questions()
    print(f"chunk={n_chunk} source={n_source} question={n_q}")


if __name__ == "__main__":
    asyncio.run(main())
